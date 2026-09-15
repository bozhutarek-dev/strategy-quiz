#!/usr/bin/env python3
"""从 GitHub 仓库 results/ 目录拉取全部成绩记录，汇总导出为 Excel。

用法：
    python3 export_results.py [输出文件名.xlsx]

依赖：本机已登录 gh CLI（gh auth login）；openpyxl（Kimi 运行时自带）。
"""
import base64
import subprocess
import sys
from datetime import datetime

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OWNER, REPO = 'bozhutarek-dev', 'strategy-quiz'
API = f'https://api.github.com/repos/{OWNER}/{REPO}'


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else '成绩汇总.xlsx'
    token = subprocess.check_output(['gh', 'auth', 'token'], text=True).strip()
    h = {'Authorization': f'token {token}', 'Accept': 'application/vnd.github+json'}

    r = requests.get(f'{API}/contents/results', headers=h)
    if r.status_code == 404:
        print('仓库中还没有任何成绩记录（results/ 目录不存在）。')
        return
    r.raise_for_status()

    records = []
    for item in r.json():
        if not item['name'].endswith('.json'):
            continue
        f = requests.get(item['url'], headers=h)
        f.raise_for_status()
        import json
        records.append(json.loads(base64.b64decode(f.json()['content']).decode('utf-8')))

    if not records:
        print('results/ 目录中没有成绩记录。')
        return

    records.sort(key=lambda x: x.get('submittedAt', ''))

    wb = Workbook()
    ws = wb.active
    ws.title = '成绩汇总'

    header = ['姓名', '得分', '总分', '正确率(%)', '提交时间'] + [f'Q{i}' for i in range(1, 21)]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E1F2')
        cell.alignment = Alignment(horizontal='center')

    for rec in records:
        ts = rec.get('submittedAt', '')
        try:
            ts = datetime.fromisoformat(ts.replace('Z', '+00:00')).astimezone().strftime('%Y-%m-%d %H:%M:%S')
        except Exception:
            pass
        marks = {a['id']: ('✓' if a['correct'] else '✗') for a in rec.get('answers', [])}
        row = [rec.get('name', ''), rec.get('score', ''), rec.get('total', ''),
               rec.get('percent', ''), ts] + [marks.get(i, '') for i in range(1, 21)]
        ws.append(row)

    for col, width in zip('ABCDE', [14, 8, 8, 11, 20]):
        ws.column_dimensions[col].width = width
    for i in range(6, 26):
        ws.column_dimensions[chr(64 + i) if i <= 26 else 'Z'].width = 5

    # 标红答错的题
    red = Font(color='CC0000')
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=25):
        for cell in row:
            if cell.value == '✗':
                cell.font = red

    wb.save(out)
    print(f'已导出 {len(records)} 条成绩记录 → {out}')


if __name__ == '__main__':
    main()
